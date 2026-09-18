"""
Multi-format document parser supporting PDF, DOCX, TXT, and scanned images via OCR.
Supports English, Hindi, and Kannada OCR via Tesseract.
"""

import io
import os
import shutil
from pathlib import Path
from typing import Union, BinaryIO, Optional
from PIL import Image

from .cleaner import clean_legal_text


def is_tesseract_available() -> bool:
    """Check if tesseract binary is accessible in system PATH or common locations."""
    if shutil.which("tesseract"):
        return True

    # Windows common locations
    win_paths = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        os.path.expandvars(r"%LOCALAPPDATA%\Tesseract-OCR\tesseract.exe"),
    ]
    for p in win_paths:
        if os.path.isfile(p):
            import pytesseract
            pytesseract.pytesseract.tesseract_cmd = p
            return True

    return False


def extract_text_from_pdf(file_input: Union[str, Path, BinaryIO, bytes]) -> str:
    """Extract text from PDF using pypdf."""
    from pypdf import PdfReader

    if isinstance(file_input, (str, Path)):
        reader = PdfReader(str(file_input))
    elif isinstance(file_input, bytes):
        reader = PdfReader(io.BytesIO(file_input))
    else:
        reader = PdfReader(file_input)

    extracted_pages = []
    for idx, page in enumerate(reader.pages):
        page_text = page.extract_text()
        if page_text:
            extracted_pages.append(page_text)

    full_text = "\n\n".join(extracted_pages)
    return clean_legal_text(full_text)


def extract_text_from_docx(file_input: Union[str, Path, BinaryIO, bytes]) -> str:
    """Extract text from DOCX using python-docx."""
    import docx

    if isinstance(file_input, bytes):
        doc = docx.Document(io.BytesIO(file_input))
    elif isinstance(file_input, (str, Path)):
        doc = docx.Document(str(file_input))
    else:
        doc = docx.Document(file_input)

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    return clean_legal_text("\n\n".join(paragraphs))


def extract_text_from_image(
    file_input: Union[str, Path, BinaryIO, bytes],
    lang: str = "eng+hin+kan"
) -> str:
    """
    Extract text from an image using Tesseract OCR (with eng, hin, kan language packs).
    Gracefully handles environment where Tesseract binary is not installed locally.
    """
    import pytesseract

    if not is_tesseract_available():
        raise RuntimeError(
            "Tesseract OCR is not installed or not in PATH on this machine. "
            "Please install Tesseract OCR with Kannada and Hindi language packs. "
            "See docs/DEPLOYMENT.md for OS-specific install commands."
        )

    if isinstance(file_input, bytes):
        image = Image.open(io.BytesIO(file_input))
    elif isinstance(file_input, (str, Path)):
        image = Image.open(str(file_input))
    else:
        image = Image.open(file_input)

    # Convert to RGB if palette or alpha
    if image.mode not in ("L", "RGB"):
        image = image.convert("RGB")

    try:
        text = pytesseract.image_to_string(image, lang=lang)
    except pytesseract.TesseractError:
        # Fall back to English if Indic language packs are missing
        text = pytesseract.image_to_string(image, lang="eng")

    return clean_legal_text(text)


def extract_text_from_file(
    file_input: Union[str, Path, BinaryIO, bytes],
    filename: Optional[str] = None
) -> str:
    """
    Extract clean text from supported document types:
    .pdf, .docx, .txt, .png, .jpg, .jpeg
    """
    # Determine extension
    ext = ""
    if filename:
        ext = Path(filename).suffix.lower()
    elif isinstance(file_input, (str, Path)):
        ext = Path(file_input).suffix.lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_input)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_input)
    elif ext in [".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp"]:
        return extract_text_from_image(file_input)
    elif ext in [".txt", ".md", ".json"]:
        if isinstance(file_input, (str, Path)):
            with open(file_input, "r", encoding="utf-8", errors="replace") as f:
                return clean_legal_text(f.read())
        elif isinstance(file_input, bytes):
            return clean_legal_text(file_input.decode("utf-8", errors="replace"))
        else:
            content = file_input.read()
            if isinstance(content, bytes):
                return clean_legal_text(content.decode("utf-8", errors="replace"))
            return clean_legal_text(content)
    else:
        # Try reading as plain text
        if isinstance(file_input, bytes):
            return clean_legal_text(file_input.decode("utf-8", errors="replace"))
        elif isinstance(file_input, (str, Path)):
            with open(file_input, "r", encoding="utf-8", errors="replace") as f:
                return clean_legal_text(f.read())
        else:
            content = file_input.read()
            if isinstance(content, bytes):
                return clean_legal_text(content.decode("utf-8", errors="replace"))
            return clean_legal_text(str(content))
